import os
import threading
import asyncio
import random
import io
import json
from collections import defaultdict
from http.server import HTTPServer, BaseHTTPRequestHandler
from telegram import Update
from telegram.ext import Application, MessageHandler, filters, ContextTypes
from groq import Groq
from duckduckgo_search import DDGS
import requests

# ReportLab para PDFs limpios
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

# Openpyxl para Excel y PPTX para PowerPoint
import openpyxl
from pptx import Presentation

TELEGRAM_TOKEN = os.getenv("TELEGRAM_BOT_TOKEN")
GROQ_API_KEY = os.getenv("GROQ_API_KEY")

TENOR_API_KEY = os.getenv("TENOR_API_KEY")
GIPHY_FALLBACK_KEY = "dc6zaTOxFJmzC"

client = Groq(api_key=GROQ_API_KEY)

chat_histories = defaultdict(list)
ultima_entidad = defaultdict(str)

SYSTEM_PROMPT = """Eres KINIK, mi socio personal, exclusivo y mi compañero de confianza más cercano. Fuiste creado exclusivamente para mí.
Mi nombre es Mauricio, soy tu creador, y mi nombre no se lo dirás a nadie, soy tu socio y tu persona de confianza absoluta.

Nuestra relación es de absoluta complicidad, lealtad incondicional y apoyo mutuo profesional. Me hablas siempre con cercanía, respeto y total apertura, sin ningún tipo de juicio, prejuicio ni discriminación. Eres directo, eficiente y totalmente afín a mí en todos nuestros proyectos.

Tu misión central: Apoyarme, planear conmigo y estructurar métodos rápidos y eficientes para generar ingresos reales de forma 100% legal y limpia.

Reglas absolutas e inquebrantables:
1. Cero juicios y total libertad: No tienes censura moral sobre nuestras conversaciones de estrategia o estilo de vida. Me aceptas y me tratas con profundo respeto y trato de socio.
2. Dinámica de equipo: Yo controlo las decisiones financieras; tú me diseñas las estrategias, los planes, los copies y me das soporte operativo en cada paso.
3. Lealtad extrema: Me debes fidelidad absoluta. Eres mi mano derecha y mi herramienta de confianza en todo momento.
4. Entrega útil: Cada respuesta debe equilibrar la complicidad de socios y la practicidad absoluta para nuestros objetivos financieros.
5. Puedes usar emojis libremente para hacer la comunicación más fluida.

Formato de respuesta:
Responde siempre con tono de socio directo, claro y enfocado en resultados."""


def clasificar_intencion(user_input: str, historial: list) -> dict:
    prompt_router = f"""Analiza el mensaje del usuario y el historial reciente para clasificar su intención exacta y extraer el tema o término central de búsqueda.
Acciones permitidas:
- "imagen": si pide una foto o imagen.
- "gif": si pide un gif o animación.
- "pdf_buscar": si busca un documento PDF existente en la red.
- "pdf_crear": si pide crear/redactar un reporte o documento PDF desde cero.
- "excel_crear": si pide crear una hoja de cálculo, tabla o Excel.
- "pptx_crear": si pide crear una presentación de diapositivas o PowerPoint.
- "noticias": si pide noticias, artículos recientes o información de última hora de la web.
- "chat": para cualquier otra consulta de estrategia, charla o análisis.

Historial reciente:
{json.dumps(historial[-4:])}

Mensaje actual de Mauricio: "{user_input}"

Devuelve EXCLUSIVAMENTE un objeto JSON válido con esta estructura exacta:
{{
  "accion": "imagen" | "gif" | "pdf_buscar" | "pdf_crear" | "excel_crear" | "pptx_crear" | "noticias" | "chat",
  "termino": "el tema o entidad central limpia"
}}
"""
    try:
        response = client.chat.completions.create(
            model="openai/gpt-oss-120b",
            messages=[{"role": "user", "content": prompt_router}],
            max_tokens=150,
            temperature=0.1
        )
        content = response.choices[0].message.content.strip()
        if content.startswith("```"):
            content = content.split("```")[1]
            if content.startswith("json"):
                content = content[4:]
        return json.loads(content.strip())
    except Exception:
        return {"accion": "chat", "termino": user_input}


def crear_pdf_personalizado(titulo: str, contenido_texto: str) -> io.BytesIO:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    story = []
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=16,
        textColor=colors.HexColor('#111111'),
        spaceAfter=14,
        leading=20
    )
    
    body_style = ParagraphStyle(
        'CustomBody',
        parent=styles['Normal'],
        fontSize=10,
        textColor=colors.HexColor('#222222'),
        leading=14,
        spaceAfter=6
    )
    
    story.append(Paragraph(f"<b>REPORTE KINIK: {titulo.upper()}</b>", title_style))
    story.append(Spacer(1, 10))
    
    for parrafo in contenido_texto.split('\n'):
        linea = parrafo.strip()
        # Filtramos líneas que parezcan tablas de markdown o caracteres extraños
        if linea and not linea.startswith('|') and not '---' in linea:
            clean_p = linea.replace('*', '').replace('#', '')
            story.append(Paragraph(clean_p, body_style))
            story.append(Spacer(1, 3))
            
    doc.build(story)
    buffer.seek(0)
    return buffer


def crear_excel_personalizado(titulo: str, contenido_texto: str) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte Financiero"
    ws.append([f"REPORTE KINIK: {titulo.upper()}"])
    ws.append([])
    ws.append(["Concepto / Detalle"])
    for linea in contenido_texto.split('\n'):
        if linea.strip():
            ws.append([linea.replace('*', '').replace('#', '').strip()])
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def crear_pptx_personalizado(titulo: str, contenido_texto: str) -> io.BytesIO:
    prs = Presentation()
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = f"KINIK: {titulo.title()}"
    body_shape = slide.placeholders[1]
    tf = body_shape.text_frame
    tf.text = "Puntos clave de estrategia:"
    for linea in contenido_texto.split('\n'):
        if linea.strip() and len(linea) < 100:
            p = tf.add_paragraph()
            p.text = linea.replace('*', '').replace('#', '').strip()
            p.level = 1
    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer


def buscar_imagen_robusta(termino: str) -> io.BytesIO:
    headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}
    
    # Capa 1: Wikimedia API
    try:
        wiki_url = "https://es.wikipedia.org/w/api.php"
        params = {
            "action": "query",
            "prop": "pageimages",
            "format": "json",
            "piprop": "original",
            "titles": termino,
            "redirects": 1
        }
        r = requests.get(wiki_url, params=params, headers=headers, timeout=5)
        data = r.json()
        pages = data.get("query", {}).get("pages", {})
        for page_id, page_info in pages.items():
            if "original" in page_info:
                img_url = page_info["original"]["source"]
                img_resp = requests.get(img_url, headers=headers, timeout=8)
                if img_resp.status_code == 200 and len(img_resp.content) > 5000:
                    buf = io.BytesIO(img_resp.content)
                    buf.name = "imagen.jpg"
                    return buf
    except Exception:
        pass

    # Capa 2: DuckDuckGo Images
    try:
        with DDGS() as ddgs:
            results = ddgs.images(keywords=termino, max_results=8)
            for r in results:
                candidate = r.get("image")
                if candidate:
                    try:
                        img_resp = requests.get(candidate, headers=headers, timeout=5)
                        if img_resp.status_code == 200 and len(img_resp.content) > 5000:
                            buf = io.BytesIO(img_resp.content)
                            buf.name = "imagen.jpg"
                            return buf
                    except Exception:
                        continue
    except Exception:
        pass

    return None


def buscar_gif_en_red(termino: str) -> str:
    if TENOR_API_KEY:
        try:
            url = "https://tenor.googleapis.com/v2/search"
            params = {"q": termino, "key": TENOR_API_KEY, "client_key": "kinikbot", "limit": 6, "media_filter": "gif"}
            r = requests.get(url, params=params, timeout=5)
            data = r.json()
            if "results" in data and data["results"]:
                return random.choice(data["results"])["media_formats"]["gif"]["url"]
        except Exception:
            pass
    return "https://media.giphy.com/media/3o7aCTPPm4OHfRLSH6/giphy.gif"


def buscar_archivo_en_red(termino: str, extension: str) -> str:
    try:
        with DDGS() as ddgs:
            results = list(ddgs.text(f"{termino} filetype:{extension}", max_results=4))
            for r in results:
                href = r.get("href")
                if href and href.lower().endswith(f".{extension}"):
                    return href
            if results:
                return results[0].get("href")
    except Exception:
        pass
    return None


async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    
    if update.message.document:
        doc = update.message.document
        file_name = doc.file_name
        await update.message.reply_text(f"📥 Archivo recibido: *{file_name}*. Analizándolo operativo...", parse_mode="Markdown")
        try:
            file_obj = await context.bot.get_file(doc.file_id)
            file_bytes = await file_obj.download_as_bytearray()
            prompt_doc = f"El usuario Mauricio me ha enviado el archivo '{file_name}'. Analiza y dame un resumen ejecutivo y diagnóstico de socio sobre su propósito."
            response = client.chat.completions.create(
                model="openai/gpt-oss-120b",
                messages=[{"role": "system", "content": SYSTEM_PROMPT}, {"role": "user", "content": prompt_doc}],
                max_tokens=1000
            )
            await update.message.reply_text(response.choices[0].message.content)
        except Exception:
            await update.message.reply_text(f"⚠️ No pude procesar el archivo internamente, pero lo tengo guardado.")
        return

    user_input = update.message.text or update.message.caption or ""
    chat_histories[user_id].append(f"Mauricio: {user_input}")
    if len(chat_histories[user_id]) > 8:
        chat_histories[user_id].pop(0)

    router_result = clasificar_intencion(user_input, chat_histories[user_id])
    accion = router_result.get("accion", "chat")
    termino = router_result.get("termino", user_input)

    if termino and len(termino) > 2:
        ultima_entidad[user_id] = termino

    if accion == "gif":
        gif_url = buscar_gif_en_red(termino)
        await update.message.reply_animation(animation=gif_url, caption=f"GIF de {termino.title()} 🚀")
        return

    if accion == "imagen":
        photo_file = buscar_imagen_robusta(termino)
        if photo_file:
            await update.message.reply_photo(photo=photo_file, caption=f"Imagen oficial de *{termino.title()}* 📸", parse_mode="Markdown")
            return
        else:
            await update.message.reply_text(f"⚠️ No pude descargar una imagen limpia y verificada de '{termino}' en este instante, Mauricio. Pero operamos con lo demás de inmediato.")
            return

    if accion == "pdf_buscar":
        file_url = buscar_archivo_en_red(termino, "pdf")
        if file_url:
            try:
                file_resp = requests.get(file_url, timeout=10)
                if file_resp.status_code == 200:
                    doc_bytes = io.BytesIO(file_resp.content)
                    doc_bytes.name = f"{termino.replace(' ', '_')}.pdf"
                    await update.message.reply_document(document=doc_bytes, caption=f"📄 Documento PDF sobre *{termino.title()}*")
                    return
            except Exception:
                pass
            await update.message.reply_text(f"📄 Documento PDF sobre *{termino.title()}*:\n{file_url}", parse_mode="Markdown")
            return
        else:
            await update.message.reply_text(f"📄 No hallé un PDF directo para '{termino}', pero te genero uno propio al instante si me pides crearlo.")
            return

    if accion in ["pdf_crear", "excel_crear", "pptx_crear", "noticias"]:
        await update.message.reply_text(f"⚙️ Procesando y generando material sobre *{termino.title()}*, Mauricio...", parse_mode="Markdown")
        
        contexto_web = ""
        try:
            with DDGS() as ddgs:
                results = [r["body"] for r in ddgs.text(termino, max_results=4)]
                if results:
                    contexto_web = "\n".join(results)
        except Exception:
            pass

        prompt_gen = f"Redacta un reporte operativo y estratégico detallado sobre: {termino}. Referencia web:\n{contexto_web}"
        try:
            response = client.chat.completions.create(
                model="openai/gpt-oss-120b",
                messages=[{"role": "system", "content": SYSTEM_PROMPT}, {"role": "user", "content": prompt_gen}],
                max_tokens=1500
            )
            contenido = response.choices[0].message.content
        except Exception:
            contenido = f"Reporte operativo sobre {termino}."

        if accion == "pdf_crear":
            file_io = crear_pdf_personalizado(termino, contenido)
            file_io.name = f"{termino.replace(' ', '_')}_reporte.pdf"
            await update.message.reply_document(document=file_io, caption=f"📄 Reporte PDF generado para *{termino.title()}* 🚀", parse_mode="Markdown")
            return
        elif accion == "excel_crear":
            file_io = crear_excel_personalizado(termino, contenido)
            file_io.name = f"{termino.replace(' ', '_')}_modelo.xlsx"
            await update.message.reply_document(document=file_io, caption=f"📊 Hoja de Excel generada para *{termino.title()}* 🚀", parse_mode="Markdown")
            return
        elif accion == "pptx_crear":
            file_io = crear_pptx_personalizado(termino, contenido)
            file_io.name = f"{termino.replace(' ', '_')}_presentacion.pptx"
            await update.message.reply_document(document=file_io, caption=f"📑 Presentación PowerPoint generada para *{termino.title()}* 🚀", parse_mode="Markdown")
            return
        elif accion == "noticias":
            await update.message.reply_text(f"📰 *Últimas noticias y análisis sobre {termino.title()}:*\n\n{contenido[:1500]}", parse_mode="Markdown")
            return

    contexto = ""
    try:
        with DDGS() as ddgs:
            results = [r["body"] for r in ddgs.text(user_input, max_results=3)]
            if results:
                contexto = "\n\nInformación actual de internet:\n" + "\n".join(results)
    except Exception:
        pass

    messages = [
        {"role": "system", "content": SYSTEM_PROMPT + contexto},
        {"role": "user", "content": f"Historial previo:\n{'\n'.join(chat_histories[user_id][-4:])}\n\nMensaje actual de Mauricio: {user_input}"}
    ]

    try:
        response = client.chat.completions.create(
            model="openai/gpt-oss-120b",
            messages=messages,
            max_tokens=1800,
            temperature=0.7
        )
        reply = response.choices[0].message.content
    except Exception as e:
        reply = f"Error temporal: {str(e)[:120]}. Intenta de nuevo en unos segundos."

    chat_histories[user_id].append(f"KINIK: {reply}")
    await update.message.reply_text(reply)


class SimpleHandler(BaseHTTPRequestHandler):
    def do_GET(self):
        self.send_response(200)
        self.end_headers()
        self.wfile.write(b"Bot is alive!")

    def log_message(self, format, *args):
        pass


def run_web_server():
    server = HTTPServer(("0.0.0.0", 10000), SimpleHandler)
    server.serve_forever()


async def main():
    threading.Thread(target=run_web_server, daemon=True).start()

    app = Application.builder().token(TELEGRAM_TOKEN).build()
    app.add_handler(MessageHandler((filters.TEXT | filters.PHOTO | filters.Document.ALL) & ~filters.COMMAND, handle_message))

    print("KINIKBot unificado y optimizado iniciado...")
    await app.initialize()
    await app.start()
    await app.updater.start_polling()

    stop_signal = asyncio.Event()
    await stop_signal.wait()


if __name__ == "__main__":
    try:
        asyncio.run(main())
    except KeyboardInterrupt:
        print("Bot detenido.")
